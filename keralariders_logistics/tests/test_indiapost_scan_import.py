"""Seller-portal spreadsheet: one India Post scan adjustment per shipment.

The fixture workbook uses the live export's columns (including the billed
amount spelled ``tarrif``). Rows are synthetic. No India Post HTTP.
"""

import io
import os
from unittest.mock import patch

from odoo.tests import TransactionCase, tagged

from odoo.addons.keralariders_logistics.tests.test_indiapost_scan_adjustment import (
    ScanAdjustmentMixin,
)

FIXTURE = os.path.join(
    os.path.dirname(__file__), 'fixtures', 'indiapost_seller_portal_scan.xlsx')

ARTICLE_DEBIT = 'ZZ900000001IN'
ARTICLE_ZERO = 'ZZ900000002IN'
ARTICLE_UNKNOWN = 'ZZ900000099IN'


@tagged('post_install', '-at_install')
class TestIndiapostScanImport(ScanAdjustmentMixin, TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._ip_make_hermetic(enabled=True)
        cls.seller = cls.env['logistics.seller'].create({
            'name': 'Scan Import Seller',
            'email': 'scan.import.seller@example.com',
            'zip': '682001',
        })
        cls.seller.write({'fulfilment_method': 'indiapost'})
        cls.wallet = cls.seller.wallet_ids[0]
        with open(FIXTURE, 'rb') as handle:
            cls.fixture = handle.read()

    def setUp(self):
        super().setUp()
        self.env['logistics.wallet.transaction'].create({
            'wallet_id': self.wallet.id,
            'amount': 5000.0,
            'reference': 'Test top-up',
        })
        self.wallet.invalidate_recordset(['balance'])
        self.order = self.env['logistics.order'].create({
            'seller_id': self.seller.id,
        })
        self.debit_shipment = self._new_shipment(
            ARTICLE_DEBIT,
            order_payment_type='cod',
            total_order_value=299.0,
            cod_amount=299.0,
        )
        self.zero_shipment = self._new_shipment(ARTICLE_ZERO)

    def _new_shipment(self, article, **overrides):
        vals = {
            'order_id': self.order.id,
            'seller_id': self.seller.id,
            'shipping_to_name': 'Scan Import Customer',
            'shipping_to_address': '12 Test Road, Test Nagar',
            'shipping_to_zip': '695001',
            'shipping_to_mobile': '9876543210',
            'item_description': 'Scan import article',
            'total_weight': 1.5,
            'length_cm': 30,
            'breadth_cm': 20,
            'height_cm': 15,
        }
        vals.update(overrides)
        shipment = self.env['logistics.shipment'].create(vals)
        shipment.sudo().with_context(allow_shipment_state_write=True).write({
            'indiapost_article_number': article,
            'state': 'in_transit',
        })
        return shipment

    def _import(self, raw, filename='indiapost_seller_portal_scan.xlsx'):
        return self.env['logistics.indiapost.scan.import'].create_from_xlsx(
            raw, filename)

    def _line(self, scan_import, article):
        return scan_import.line_ids.filtered(
            lambda line: line.article_number == article)

    def _xlsx_rows(self, data_rows, extra_headers=()):
        """Build a workbook with the fixture's headers plus any extra columns."""
        import openpyxl
        base = openpyxl.load_workbook(io.BytesIO(self.fixture), read_only=True)
        try:
            headers = list(next(base.active.iter_rows(values_only=True)))
        finally:
            base.close()
        headers = [header for header in headers if header] + list(extra_headers)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Data'
        sheet.append(headers)
        for row in data_rows:
            sheet.append([row.get(header, '') for header in headers])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_fixture_column_map(self):
        rows = self.env['logistics.indiapost.scan.import']._ip_parse_scan_xlsx(
            self.fixture)
        self.assertEqual(
            [row['article'] for row in rows],
            [ARTICLE_DEBIT, ARTICLE_ZERO, ARTICLE_UNKNOWN],
        )
        self.assertAlmostEqual(rows[0]['billed'], 150.0)
        self.assertAlmostEqual(rows[0]['base_amount'], 100.0)
        self.assertEqual(rows[0]['weight_g'], 800)
        self.assertEqual(rows[0]['booking_ref'], '26090001')
        self.assertEqual(rows[0]['office_name'], 'Test SO')
        self.assertEqual(rows[0]['booked_on'], '2026/10/01T10:00:00')
        self.assertFalse(rows[0]['length_cm'])

    def test_import_posts_one_difference_and_unknown_is_an_error(self):
        self._debit_pickup(self.debit_shipment, total=118.0)
        self._debit_pickup(self.zero_shipment, total=80.0)

        scan_import = self._import(self.fixture)

        self.assertEqual(scan_import.row_count, 3)
        self.assertEqual(scan_import.posted_count, 1)
        self.assertEqual(scan_import.skipped_count, 1)
        self.assertEqual(scan_import.error_count, 1)
        self.assertAlmostEqual(scan_import.total_debit, 32.0)
        self.assertAlmostEqual(scan_import.total_credit, 0.0)

        posted = self._line(scan_import, ARTICLE_DEBIT)
        self.assertEqual(posted.state, 'posted')
        self.assertEqual(posted.shipment_id, self.debit_shipment)
        self.assertEqual(posted.seller_id, self.seller)
        self.assertAlmostEqual(posted.previous_charge, 118.0)
        self.assertAlmostEqual(posted.file_amount, 150.0)
        self.assertAlmostEqual(posted.base_amount, 100.0)
        self.assertAlmostEqual(posted.difference, 32.0)
        self.assertAlmostEqual(posted.wallet_amount, -32.0)
        self.assertEqual(posted.amount_source, 'portal_bill')
        wallet_lines = self._adj_lines(self.debit_shipment)
        self.assertEqual(len(wallet_lines), 1)
        self.assertEqual(wallet_lines, posted.wallet_txn_id)
        self.assertAlmostEqual(wallet_lines.amount, -32.0)
        self.assertEqual(
            wallet_lines.reference,
            'IP-SCAN-ADJ:%s' % self.debit_shipment.id,
        )
        self.assertTrue(self.debit_shipment.indiapost_scan_adjusted)
        self.assertAlmostEqual(self.debit_shipment.indiapost_scan_quote, 150.0)
        self.assertAlmostEqual(
            self.debit_shipment.indiapost_scan_difference, 32.0)
        self.assertEqual(
            self.debit_shipment.indiapost_scan_quote_source, 'portal_bill')
        self.assertEqual(self.debit_shipment.indiapost_scan_weight_g, 800)
        self.assertEqual(self.debit_shipment.cod_amount, 299.0)
        self.assertFalse(self.env['logistics.account.transfer'].search([
            ('shipment_id', '=', self.debit_shipment.id),
            ('transfer_type', '=', 'cod_payment'),
        ]))
        mail = self.env['mail.mail'].search([
            ('res_id', '=', self.debit_shipment.id),
            ('model', '=', 'logistics.shipment'),
            ('subject', 'ilike', 'India Post updated'),
        ])
        self.assertEqual(len(mail), 1)

        zero = self._line(scan_import, ARTICLE_ZERO)
        self.assertEqual(zero.state, 'skipped')
        self.assertEqual(zero.reason, 'zero difference')
        self.assertFalse(self._adj_lines(self.zero_shipment))
        self.assertFalse(self.zero_shipment.indiapost_scan_adjusted)

        unknown = self._line(scan_import, ARTICLE_UNKNOWN)
        self.assertEqual(unknown.state, 'error')
        self.assertEqual(unknown.reason, 'article not found')
        self.assertFalse(unknown.shipment_id)

    def test_reimport_does_not_post_a_second_wallet_line(self):
        self._debit_pickup(self.debit_shipment, total=118.0)
        self._debit_pickup(self.zero_shipment, total=80.0)
        self._import(self.fixture)

        again = self._import(self.fixture)

        self.assertEqual(len(self._adj_lines(self.debit_shipment)), 1)
        self.assertAlmostEqual(
            self._adj_lines(self.debit_shipment).amount, -32.0)
        skipped = self._line(again, ARTICLE_DEBIT)
        self.assertEqual(skipped.state, 'skipped')
        self.assertEqual(skipped.reason, 'already adjusted')
        self.assertEqual(again.posted_count, 0)
        unknown = self._line(again, ARTICLE_UNKNOWN)
        self.assertEqual(unknown.state, 'error')
        self.assertEqual(unknown.reason, 'article not found')

    def test_same_file_does_not_post_twice(self):
        self._debit_pickup(self.debit_shipment, total=118.0)
        raw = self._xlsx_rows([
            {
                'article-number': ARTICLE_DEBIT,
                'tarrif': '150',
                'weight': '800',
                'customer-bulk-reference': '26090001',
            },
            {
                'article-number': ARTICLE_DEBIT,
                'tarrif': '180',
                'weight': '900',
                'customer-bulk-reference': '26090001',
            },
        ])
        scan_import = self._import(raw)
        self.assertEqual(scan_import.posted_count, 1)
        self.assertEqual(scan_import.skipped_count, 1)
        self.assertEqual(len(self._adj_lines(self.debit_shipment)), 1)
        self.assertAlmostEqual(
            self._adj_lines(self.debit_shipment).amount, -32.0)
        skipped = scan_import.line_ids.filtered(lambda line: line.state == 'skipped')
        self.assertEqual(skipped.reason, 'already adjusted')

    def test_zero_difference_can_still_be_adjusted_later(self):
        self._debit_pickup(self.zero_shipment, total=80.0)
        first = self._import(self.fixture)
        self.assertEqual(self._line(first, ARTICLE_ZERO).reason, 'zero difference')
        self.assertFalse(self.zero_shipment.indiapost_scan_adjusted)

        raw = self._xlsx_rows([{
            'article-number': ARTICLE_ZERO,
            'tarrif': '95',
            'weight': '450',
            'customer-bulk-reference': '26090002',
        }])
        second = self._import(raw)
        posted = second.line_ids
        self.assertEqual(posted.state, 'posted')
        self.assertAlmostEqual(posted.difference, 15.0)
        self.assertAlmostEqual(posted.wallet_amount, -15.0)
        self.assertEqual(len(self._adj_lines(self.zero_shipment)), 1)
        self.assertTrue(self.zero_shipment.indiapost_scan_adjusted)

    def test_flag_without_wallet_line_is_skipped(self):
        self._debit_pickup(self.debit_shipment, total=118.0)
        self.debit_shipment.with_context(
            allow_delivery_charge_write=True).write({
                'indiapost_scan_adjusted': True,
            })
        scan_import = self._import(self.fixture)
        skipped = self._line(scan_import, ARTICLE_DEBIT)
        self.assertEqual(skipped.state, 'skipped')
        self.assertEqual(skipped.reason, 'already adjusted')
        self.assertFalse(self._adj_lines(self.debit_shipment))

    def test_booking_reference_matches_awb_when_article_does_not(self):
        awb = self._unused_awb()
        shipment = self._new_shipment('ZZ900000003IN', name=awb)
        self._debit_pickup(shipment, total=100.0)
        raw = self._xlsx_rows([{
            'article-number': 'ZZ000000000IN',
            'tarrif': '110',
            'weight': '600',
            'customer-bulk-reference': awb,
        }])
        scan_import = self._import(raw)
        line = scan_import.line_ids
        self.assertEqual(line.state, 'posted')
        self.assertEqual(line.shipment_id, shipment)
        self.assertAlmostEqual(line.wallet_amount, -10.0)
        self.assertEqual(len(self._adj_lines(shipment)), 1)

    def test_file_bill_is_used_when_dimensions_are_also_present(self):
        self._debit_pickup(self.debit_shipment, total=118.0)

        def fail_quote(*args, **kwargs):
            raise AssertionError('tariff quote must not run when tarrif is set')

        raw = self._xlsx_rows(
            [{
                'article-number': ARTICLE_DEBIT,
                'tarrif': '130',
                'weight': '2000',
                'length-cm': 40,
                'breadth-cm': 30,
                'height-cm': 20,
            }],
            extra_headers=('length-cm', 'breadth-cm', 'height-cm'),
        )
        with patch.object(
                self.registry['logistics.indiapost.tariff'],
                'quote', autospec=True, side_effect=fail_quote):
            scan_import = self._import(raw)
        line = scan_import.line_ids
        self.assertEqual(line.state, 'posted')
        self.assertEqual(line.amount_source, 'portal_bill')
        self.assertAlmostEqual(line.file_amount, 130.0)
        self.assertAlmostEqual(line.wallet_amount, -12.0)

    def test_missing_bill_requotes_from_weight_and_dimensions(self):
        self._debit_pickup(self.debit_shipment, total=118.0)
        captured = {}

        def fake_quote(*args, **kwargs):
            captured.update(kwargs)
            return {
                'ok': True,
                'base_tariff': 120.0,
                'vas_charges': 0.0,
                'total_tax': 20.0,
                'final_amount': 140.0,
                'total_payable': 140.0,
                'billed_weight_g': 2000,
                'chargeable_weight_g': 4800,
                'volumetric_weight_g': 4800,
                'actual_weight_g': 2000,
            }

        raw = self._xlsx_rows(
            [{
                'article-number': ARTICLE_DEBIT,
                'tarrif': '',
                'weight': '2000',
                'length-cm': 40,
                'breadth-cm': 30,
                'height-cm': 20,
            }],
            extra_headers=('length-cm', 'breadth-cm', 'height-cm'),
        )
        with patch.object(
                self.registry['logistics.indiapost.tariff'],
                'quote', autospec=True, side_effect=fake_quote):
            scan_import = self._import(raw)
        line = scan_import.line_ids
        self.assertEqual(line.state, 'posted')
        self.assertEqual(line.amount_source, 'tariff')
        self.assertAlmostEqual(line.file_amount, 140.0)
        self.assertAlmostEqual(line.wallet_amount, -22.0)
        self.assertEqual(captured.get('weight_g'), 2000)
        self.assertEqual(captured.get('article_type'), self.debit_shipment._ip_product())
        self.assertEqual(captured.get('length_cm'), 40)
        self.assertEqual(
            self.debit_shipment.indiapost_scan_quote_source, 'tariff')

    def _unused_awb(self):
        Shipment = self.env['logistics.shipment']
        for serial in range(99000001, 99000100):
            name = str(serial)
            if not Shipment.search_count([('name', '=', name)]):
                return name
        self.fail('no free AWB name for the booking-reference test')
